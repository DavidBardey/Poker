import copy
import io
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from openpyxl import load_workbook
import app as poker

ROOT = Path(__file__).resolve().parents[1]

class PokerTests(unittest.TestCase):
    def setUp(self):
        self.original = copy.deepcopy(poker.STATE)
        self.addCleanup(self.restore)
        poker.app.config["TESTING"] = True
        self.client = poker.app.test_client()
        self.tmp = tempfile.TemporaryDirectory()
        self.exports = poker.EXPORTS_DIR
        poker.EXPORTS_DIR = Path(self.tmp.name)
        self.configure()

    def restore(self):
        poker.STATE.clear()
        poker.STATE.update(self.original)
        poker.EXPORTS_DIR = self.exports
        self.tmp.cleanup()

    def configure(self, **changes):
        data = dict(players=3, duration=60, buyin_money=50000, buyin_chips=1000,
                    max_rebuys=1, paid_places=2, winner_pct=70, second_pct=0,
                    player_1="Alice", player_2="Bob", player_3="Chloé")
        data.update(changes)
        return self.client.post("/setup", data=data)

    def kill(self, player, outcome="out"):
        return self.client.post("/api/kill", json=dict(killed=player, killer="Alice", outcome=outcome))

    def test_routes_and_existing_translations(self):
        for code in poker.LANGS:
            with self.subTest(language=code):
                self.client.get("/set-lang/" + code)
                for route in ("/", "/setup", "/clock"):
                    response = self.client.get(route)
                    self.assertEqual(response.status_code, 200)
                    self.assertIn('lang="' + code + '"', response.get_data(as_text=True))
                self.assertIn(poker.TEXTS[code]["clock"], self.client.get("/clock").get_data(as_text=True))

    def test_chip_schedule_is_independent_of_money(self):
        first = poker.STATE["schedule"].copy()
        self.configure(buyin_money=100)
        self.assertEqual(first, poker.STATE["schedule"])
        self.assertEqual(poker.total_pot(), 300)

    def test_large_stack_schedule_terminates(self):
        # Run separately so a regression cannot hang the entire suite.
        result = subprocess.run([sys.executable, "-c",
            "import app; s=app.build_schedule(100000, 12); assert len(s)==12; assert len(set(s))==12"],
            cwd=ROOT, capture_output=True, text=True, timeout=5)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_no_winner_or_unearned_payout_before_end(self):
        self.kill("Chloé")
        final = self.client.get("/api/final").json
        self.assertEqual(final["winner"], "")
        self.assertEqual(poker.ranking(), [(3, "Chloé")])
        self.assertEqual(final["payouts"], [])
        response = self.client.get("/export")
        wb = load_workbook(io.BytesIO(response.data))
        self.assertEqual(wb["Classement"]["A2"].value, 3)
        rows = list(wb["Résumé joueurs"].values)
        self.assertEqual(rows[-1][-1], "Éliminé")
        self.assertIsNone(wb["Gains"]["B8"].value)
        wb.close()
        response.close()

    def test_rebuy_elimination_pot_ranking_and_all_excel_sheets(self):
        self.assertTrue(self.kill("Bob", "rebuy").json["ok"])
        self.assertEqual(poker.total_pot(), 200000)
        self.assertFalse(self.kill("Bob", "rebuy").json["ok"])
        self.kill("Chloé")
        self.kill("Bob")
        final = self.client.get("/api/final").json
        self.assertEqual(final["winner"], "Alice")
        self.assertEqual(poker.ranking(), [(1,"Alice"),(2,"Bob"),(3,"Chloé")])
        self.assertEqual([p["amount"] for p in final["payouts"]], [140000,60000])
        response = self.client.get("/export")
        wb = load_workbook(io.BytesIO(response.data))
        self.assertEqual(wb.sheetnames, ["Journal","Classement","Gains","Résumé joueurs","Paramètres"])
        self.assertEqual(wb["Gains"]["B1"].value, 200000)
        self.assertEqual(wb["Gains"]["B3"].value, 1000)
        self.assertEqual(wb["Gains"]["D8"].value, 140000)
        self.assertEqual(wb["Résumé joueurs"]["B2"].value, 3)
        wb.close()
        response.close()

    def test_invalid_outcome_is_not_an_elimination(self):
        before = copy.deepcopy(poker.STATE)
        result = self.kill("Bob", "typo")
        self.assertFalse(result.json["ok"])
        self.assertEqual(poker.STATE, before)

    def test_kill_validation_and_name_normalization(self):
        before = copy.deepcopy(poker.STATE)
        for data in (dict(killed="Alice",killer="Alice",outcome="out"),
                     dict(killed="Missing",killer="Alice",outcome="out"),
                     dict(killed="Bob",killer="Missing",outcome="out")):
            self.assertFalse(self.client.post("/api/kill",json=data).json["ok"])
            self.assertEqual(poker.STATE,before)
        self.assertTrue(self.kill(" Bob ", "rebuy").json["ok"])
        self.assertEqual(poker.rebuys_for_player("Bob"),1)

    def test_payout_rules_one_two_three(self):
        for paid, expected in ((1,[100]),(2,[60,40]),(3,[60,30,10])):
            self.configure(paid_places=paid,winner_pct=60,second_pct=30)
            self.assertEqual(poker.payout_percentages(), expected)

    def test_standard_blind_structure_is_unchanged(self):
        self.assertEqual(poker.build_schedule(1000, 8),
                         ["10 / 20", "15 / 30", "20 / 40", "30 / 60",
                          "50 / 100", "75 / 150", "100 / 200", "150 / 300"])

    def test_live_language_endpoint_does_not_mutate_game(self):
        before = copy.deepcopy(poker.STATE)
        for code in poker.LANGS:
            response = self.client.post("/api/lang/" + code)
            self.assertEqual(response.json, poker.TEXTS[code])
            self.assertEqual(poker.STATE, before)
            with self.client.session_transaction() as session:
                self.assertEqual(session["lang"], code)
        self.assertEqual(self.client.post("/api/lang/xx").status_code, 400)

    def test_level_bounds(self):
        self.client.post("/api/level",json={"direction":"previous"})
        self.assertEqual(poker.STATE["current_level"],1)
        for _ in range(20):
            self.client.post("/api/level",json={"direction":"next"})
        self.assertEqual(poker.STATE["current_level"],len(poker.STATE["schedule"]))

    def test_unused_blank_percentages_do_not_break_setup(self):
        for paid in (1, 2):
            response = self.configure(paid_places=paid, second_pct="",
                                      winner_pct="" if paid == 1 else 70)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(self.client.get('/clock').status_code, 200)

    def test_invalid_setup_returns_localized_form_without_mutating_game(self):
        cases = [dict(winner_pct=80, second_pct=30, paid_places=3),
                 dict(player_2=" alice "), dict(second_pct="", paid_places=3),
                 dict(players=""), dict(players=21), dict(duration=0),
                 dict(buyin_money="abc"), dict(buyin_chips=0),
                 dict(max_rebuys=-1), dict(paid_places=4)]
        before = copy.deepcopy(poker.STATE)
        for code in poker.LANGS:
            self.client.get('/set-lang/' + code)
            for changes in cases:
                with self.subTest(language=code, changes=changes):
                    response = self.configure(**changes)
                    self.assertEqual(response.status_code, 400)
                    html = response.get_data(as_text=True)
                    self.assertIn('role="alert"', html)
                    self.assertIn('lang="'+code+'"', html)
                    self.assertEqual(poker.STATE, before)

    def test_1200_chips_and_money_have_separate_effects(self):
        self.assertEqual(self.configure(buyin_chips=1200, buyin_money=50).status_code, 200)
        schedule = poker.STATE['schedule'].copy()
        self.assertEqual(schedule, poker.build_schedule(1200, 8))
        self.assertEqual(poker.total_pot(), 150)
        self.configure(buyin_chips=1200, buyin_money=100)
        self.assertEqual(poker.STATE['schedule'], schedule)
        self.configure(buyin_chips=12000, buyin_money=50)
        self.assertNotEqual(poker.STATE['schedule'], schedule)
        self.assertEqual(poker.total_pot(), 150)
        self.configure(buyin_chips=1200, buyin_money=50)
        self.kill('Bob', 'rebuy')
        self.kill('Chloé')
        self.kill('Bob')
        final = self.client.get('/api/final').json
        self.assertEqual(final['pot'], 200)
        self.assertEqual([p['amount'] for p in final['payouts']], [140, 60])

    def test_clock_without_configuration_returns_to_setup(self):
        poker.STATE['configured'] = False
        response = self.client.get('/clock')
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.location.endswith('/setup'))

if __name__ == "__main__":
    unittest.main()
